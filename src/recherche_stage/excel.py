"""Tenue d'un tableau Excel de suivi des candidatures déjà existant.

Le tableau de l'utilisateur garde sa mise en page : les colonnes sont
reconnues à leur intitulé (« Entreprise », « Poste », « Date d'envoi »,
« Statut »…), les colonnes inconnues ne sont jamais touchées, et une colonne
de remarques n'est remplie que si elle est vide. Une candidature déjà présente
(même entreprise et même poste, à la normalisation près) est mise à jour ; les
autres sont ajoutées à la suite, dans le style de la dernière ligne. Quand une
colonne propose une liste déroulante, la valeur écrite est prise dans la liste.

Avant la première écriture, une copie du fichier d'origine est gardée à côté.
"""

from __future__ import annotations

import os
import re
import shutil
import unicodedata
from copy import copy
from datetime import date, datetime
from pathlib import Path

from .hachage import id_dossier
from .modeles import Dossier, Statut
from .naf import libelle as libelle_naf
from .presentiel import texte_excel

# Champ → intitulés de colonne reconnus (après normalisation : minuscules,
# sans accents ni ponctuation). Le premier intitulé égal gagne, puis le plus
# long contenu dans l'en-tête.
COLONNES: dict[str, tuple[str, ...]] = {
    "entreprise": ("entreprise", "societe", "structure", "organisme", "company", "nom de l entreprise", "employeur"),
    "poste": ("poste", "intitule", "intitule du poste", "offre", "sujet", "mission", "stage", "titre"),
    "lieu": ("lieu", "ville", "localisation", "adresse", "departement", "localite"),
    "secteur": ("secteur", "domaine", "activite"),
    "type": ("type", "type de candidature", "nature", "spontanee ou offre"),
    "canal": ("canal", "moyen", "mode", "mode de candidature", "via", "envoye par", "plateforme", "candidature via"),
    "contact": ("contact", "email", "e mail", "mail", "adresse mail", "destinataire", "recruteur", "interlocuteur"),
    "lien": ("lien", "url", "annonce", "lien annonce", "lien de l offre", "site", "source"),
    "date_candidature": ("date", "date candidature", "date de candidature", "date d envoi", "date envoi", "envoyee le",
                         "date de postulation", "postule le", "candidature le"),
    "statut": ("statut", "etat", "avancement", "situation", "etape"),
    "relance": ("relance", "relances", "date relance", "date de relance", "derniere relance", "relance le"),
    "reponse": ("reponse", "retour", "resultat", "reponse obtenue"),
    "date_reponse": ("date reponse", "date de reponse", "date retour", "reponse le"),
    "entretien": ("entretien", "date entretien", "date d entretien"),
    "score": ("score", "pertinence", "note", "interet"),
    "presentiel": ("presentiel", "depot en presentiel", "depot sur place", "sur place", "en personne", "main propre",
                   "deposer sur place", "a deposer sur place", "en main propre"),
    "remarques": ("remarques", "remarque", "commentaires", "commentaire", "notes", "observations", "infos"),
}

# Trop vagues pour être cherchés à l'intérieur d'un intitulé (« Dates du stage »,
# « Date du mail »…) : reconnus seulement quand ils sont l'intitulé entier.
EXACTS = {"stage", "mission", "titre", "type", "nature", "mode", "via", "site", "source", "note", "date",
          "infos", "etape", "mail", "offre", "activite", "situation", "retour"}

STATUTS_EXPORTES = {Statut.validee, Statut.envoyee, Statut.relancee, Statut.entretien, Statut.refus, Statut.acceptee}

LIBELLES_STATUT = {
    "a_envoyer": "À envoyer", "a_deposer": "À déposer", Statut.envoyee: "Envoyée", Statut.relancee: "Relancée",
    Statut.entretien: "Entretien", Statut.refus: "Refusée", Statut.acceptee: "Acceptée",
}
# Mots qui désignent chaque statut dans une liste déroulante, par ordre de préférence.
MOTS_STATUT = {
    "a_envoyer": ("a envoyer", "a postuler", "prepar", "brouillon", "a faire"),
    "a_deposer": ("a deposer", "a envoyer", "a postuler", "prepar", "a faire"),
    Statut.envoyee: ("envoy", "postul", "attente", "en cours", "candidat"),
    Statut.relancee: ("relanc", "attente", "en cours", "envoy"),
    Statut.entretien: ("entretien", "convoc", "en cours"),
    Statut.refus: ("refus", "non retenu", "negati", "rejet"),
    Statut.acceptee: ("accept", "retenu", "positi", "obtenu", "valide"),
}
# Démarches faites hors de l'outil (demarches.yaml) : statut → mots cherchés dans la liste, libellé sans liste.
# « En attente » ou « Contact » absents de la liste : la case reste vide, la démarche est dans les notes.
STATUTS_DEMARCHE = {
    "envoyee": (MOTS_STATUT[Statut.envoyee], "Envoyée"),
    "en_attente": (("attente", "en cours"), "En attente"),
    "a_faire": (("a faire", "a envoyer", "a postuler"), "À faire"),
    "contact": (("contact", "forum", "rencontr"), "Contact"),
}
LIBELLES_REPONSE = {"refus": "Refus", "entretien": "Entretien", "demande_info": "Demande d'informations"}


# Décrivent la candidature : sur une ligne existante, complétés seulement s'ils sont vides
# (l'utilisateur a pu écrire le nom à sa façon). Les autres suivent l'avancement.
DESCRIPTIFS = {"entreprise", "poste", "lieu", "secteur", "type", "contact", "lien", "score", "remarques", "presentiel"}

# Colonnes ajoutées au tableau quand elles y manquent et qu'il y a quelque chose à y mettre.
AJOUTABLES = {"presentiel": "Dépôt en présentiel"}
VERT = {"fill": "C6EFCE", "police": "006100"}  # style « Bon » d'Excel


class ErreurExcel(Exception):
    pass


def _norm(texte) -> str:
    texte = unicodedata.normalize("NFKD", str(texte or "")).encode("ascii", "ignore").decode().lower()
    return " ".join(re.sub(r"[^a-z0-9]+", " ", texte).split())


def reconnaitre_colonnes(entetes: dict[int, str]) -> dict[str, int]:
    """{champ: numéro de colonne} d'après les intitulés d'une ligne d'en-tête."""
    normes = {col: _norm(v) for col, v in entetes.items() if _norm(v)}
    trouvees: dict[str, int] = {}
    prises: set[int] = set()
    # 1. intitulés identiques ; 2. intitulé reconnu contenu dans l'en-tête (le plus long d'abord)
    for exact in (True, False):
        candidats = []
        for champ, noms in COLONNES.items():
            if champ in trouvees:
                continue
            for col, entete in normes.items():
                for nom in noms:
                    if entete == nom if exact else nom not in EXACTS and re.search(rf"\b{re.escape(nom)}\b", entete):
                        candidats.append((-len(nom), col, champ))
        for _, col, champ in sorted(candidats):
            if champ not in trouvees and col not in prises:
                trouvees[champ] = col
                prises.add(col)
    return trouvees


def _ligne_entete(ws) -> tuple[int, dict[str, int]]:
    meilleure = (0, {})
    for ligne in range(1, min(ws.max_row, 15) + 1):
        entetes = {c.column: c.value for c in ws[ligne] if isinstance(c.value, str)}
        colonnes = reconnaitre_colonnes(entetes)
        if len(colonnes) > len(meilleure[1]):
            meilleure = (ligne, colonnes)
    if "entreprise" not in meilleure[1]:
        raise ErreurExcel("colonne « Entreprise » introuvable dans les 15 premières lignes du tableau")
    return meilleure


def _liste_deroulante(ws, cellule) -> list[str]:
    for dv in ws.data_validations.dataValidation:
        if dv.type == "list" and cellule.coordinate in dv.sqref and dv.formula1 and dv.formula1.startswith('"'):
            return [v.strip() for v in dv.formula1.strip('"').split(",") if v.strip()]
    return []


def _choisir(valeurs: list[str], mots: tuple[str, ...], defaut: str) -> str:
    for mot in mots:
        for v in valeurs:
            n = _norm(v)
            if not re.search(rf"\b{re.escape(mot)}", n):
                continue
            if mot in ("retenu", "valide", "positi") and ("non" in n.split() or "refus" in n):
                continue
            if n.startswith("a ") and not mot.startswith("a "):
                continue  # « À envoyer » n'est pas « Envoyée » : une tâche à faire, pas un état atteint
            return v
    return defaut


# Mots sans valeur pour reconnaître une entreprise (« NORDEX FRANCE SAS » est « Nordex »).
GENERIQUES = {"sa", "sas", "sasu", "sarl", "eurl", "snc", "sca", "scop", "france", "groupe", "group", "societe",
              "et", "de", "du", "des", "la", "le", "les", "the"}
PETITS_MOTS = {"de", "du", "des", "la", "le", "les", "en", "sur", "sous", "et", "aux", "au", "d", "l", "a", "lez", "les"}


def _coeur(nom) -> str:
    """Nom réduit à l'essentiel, pour reconnaître une entreprise écrite autrement."""
    sans_parentheses = re.sub(r"\(.*?\)", " ", str(nom or ""))
    return " ".join(m for m in _norm(sans_parentheses).split() if len(m) > 1 and m not in GENERIQUES)


def nom_entreprise(nom: str) -> str:
    """« ADRE EAU (ADRE EAU) » → « ADRE EAU » ; une parenthèse utile (sigle, marque) reste."""
    m = re.match(r"^(.*?)\s*\((.*)\)\s*$", nom or "")
    if not m:
        return nom or ""
    dehors, dedans = _coeur(m.group(1)), _coeur(m.group(2))
    if dehors and dedans and (dehors == dedans or dedans.startswith(dehors) or dehors.startswith(dedans)):
        return m.group(1).strip()
    return nom


def _departement(lieu) -> str:
    m = re.search(r"\((97\d|\d[\dAB])\d*\)", str(lieu or ""))
    return m.group(1) if m else ""


def lieu_lisible(lieu: str) -> str:
    """« SAINT-QUENTIN (02100) » → « Saint-Quentin (02) », comme on l'écrit à la main."""
    m = re.match(r"^(.*?)\s*\((\d{5})\)\s*$", lieu or "")
    if not m:
        return lieu or ""
    ville, code = m.group(1).strip(), m.group(2)
    if ville == ville.upper():
        morceaux = re.split(r"([\s\-'’])", ville.lower())
        ville = "".join(x if (i and x in PETITS_MOTS) or not x.strip() else x[:1].upper() + x[1:]
                        for i, x in enumerate(morceaux))
    return f"{ville} ({code[:3] if code.startswith('97') else code[:2]})"


def _etat(d: Dossier) -> str | Statut:
    if d.statut == Statut.validee:
        return "a_envoyer"
    if d.statut == Statut.en_attente_humain and d.etape_attente == "depot":
        return "a_deposer"
    return d.statut


def a_exporter(d: Dossier) -> bool:
    return d.statut in STATUTS_EXPORTES or (d.statut == Statut.en_attente_humain and d.etape_attente == "depot")


def valeurs(d: Dossier) -> dict[str, object]:
    o = d.offre
    canal = {"email": "Mail", "formulaire": "Site de l'entreprise", "plateforme": o.source,
             "main_propre": "En main propre"}.get(d.envoi.canal or "", "")
    envoi = d.envoi.horodatage
    return {
        "entreprise": nom_entreprise(o.entreprise),
        "poste": o.intitule,
        "lieu": lieu_lisible(o.lieu),
        "secteur": libelle_naf(d.fiche_entreprise.secteur or o.secteur or ""),
        "type": "Spontanée" if o.type == "spontanee" else "Offre",
        "canal": canal,
        "contact": d.envoi.destinataire if d.envoi.canal == "email" else (o.contact_email or d.fiche_entreprise.contact or ""),
        "lien": o.url,
        "date_candidature": envoi.date() if envoi else None,
        "relance": d.suivi.relances[-1].date() if d.suivi.relances else None,
        "reponse": LIBELLES_REPONSE.get(d.suivi.reponse, ""),
        "date_reponse": d.suivi.date_reponse.date() if d.suivi.date_reponse else None,
        "entretien": d.suivi.entretien_le if d.suivi.entretien_le else None,
        "score": round(d.score.valeur) if d.score.valeur else None,
        "remarques": d.envoi.preuve if d.envoi.canal != "email" else "",
        "presentiel": texte_excel(d.presentiel),
    }


def valeurs_demarche(dm) -> dict[str, object]:
    quand = f" ({dm.date:%d/%m/%Y})" if dm.date else ""
    notes = f"{dm.canal}{quand} : {dm.notes}" if dm.canal else dm.notes
    return {
        "entreprise": dm.entreprise, "lieu": dm.lieu, "secteur": dm.activite, "canal": dm.canal,
        "contact": dm.contact, "date_candidature": dm.date if dm.statut == "envoyee" else None,
        "relance": dm.relance, "remarques": notes.strip(" :"),
    }


def _ajouter_colonne(ws, ligne_entete: int, intitule: str) -> int:
    """Ajoute une colonne après la dernière colonne titrée, dans le style de l'en-tête
    voisin, et l'intègre au tableau Excel mis en forme qui s'arrêtait juste avant."""
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import TableColumn

    derniere = max((c.column for c in ws[ligne_entete] if c.value not in (None, "")), default=0)
    col = derniere + 1
    cellule = ws.cell(ligne_entete, col, intitule)
    if derniere and ws.cell(ligne_entete, derniere).has_style:
        cellule._style = copy(ws.cell(ligne_entete, derniere)._style)
    lettre = get_column_letter(col)
    precedente = get_column_letter(derniere) if derniere else None
    if precedente and ws.column_dimensions[precedente].width:
        ws.column_dimensions[lettre].width = max(ws.column_dimensions[precedente].width, 28)
    else:
        ws.column_dimensions[lettre].width = 40
    for table in ws.tables.values():
        debut, fin = table.ref.split(":")
        if re.match(r"[A-Z]+", fin).group(0) == precedente and int(re.search(r"\d+", debut).group(0)) == ligne_entete:
            fin_ligne = re.search(r"\d+", fin).group(0)
            table.ref = f"{debut}:{lettre}{fin_ligne}"
            table.tableColumns.append(TableColumn(id=len(table.tableColumns) + 1, name=intitule))
            if table.autoFilter is not None:
                table.autoFilter.ref = table.ref
    return col


def _surligner(cellule) -> None:
    from openpyxl.styles import Font, PatternFill

    cellule.fill = PatternFill("solid", start_color=VERT["fill"], end_color=VERT["fill"])
    cellule.font = Font(color=VERT["police"], bold=True)


def mettre_a_jour(chemin: str | Path, dossiers: list[Dossier], feuille: str | None = None,
                  demarches: list = ()) -> tuple[int, int]:
    """Met le tableau à jour ; renvoie (lignes ajoutées, lignes mises à jour). Les
    démarches faites hors de l'outil ont chacune leur ligne ; sur une ligne
    existante, elles ne remplissent que les cases vides."""
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise ErreurExcel("module openpyxl absent : pip install openpyxl") from e

    chemin = Path(chemin)
    if not chemin.exists():
        raise ErreurExcel(f"fichier introuvable : {chemin}")
    sauvegarde = chemin.with_name(f"{chemin.stem}.avant_rs{chemin.suffix}")
    if not sauvegarde.exists():
        shutil.copy2(chemin, sauvegarde)

    classeur = load_workbook(chemin)
    ws = classeur[feuille] if feuille else classeur.worksheets[0]
    ligne_entete, colonnes = _ligne_entete(ws)
    a_ecrire = [d for d in dossiers if a_exporter(d)]
    for champ, intitule in AJOUTABLES.items():
        if champ not in colonnes and any(valeurs(d)[champ] for d in a_ecrire):
            colonnes[champ] = _ajouter_colonne(ws, ligne_entete, intitule)

    def cle(entreprise, poste) -> str:
        return id_dossier(str(entreprise or ""), str(poste or "")) if "poste" in colonnes else _norm(entreprise)

    existantes: dict[str, int] = {}
    coeurs: dict[tuple[str, str], int] = {}  # (nom réduit, département) -> ligne, pour les noms écrits autrement
    homonymes: dict[str, list[tuple[int, str]]] = {}  # nom réduit -> lignes, pour signaler une même entreprise
    derniere = ligne_entete
    for ligne in range(ligne_entete + 1, ws.max_row + 1):
        entreprise = ws.cell(ligne, colonnes["entreprise"]).value
        if entreprise in (None, ""):
            continue
        derniere = ligne
        poste = ws.cell(ligne, colonnes["poste"]).value if "poste" in colonnes else ""
        existantes.setdefault(cle(entreprise, poste), ligne)
        lieu = ws.cell(ligne, colonnes["lieu"]).value if "lieu" in colonnes else ""
        coeurs.setdefault((_coeur(entreprise), _departement(lieu)), ligne)
        homonymes.setdefault(_coeur(entreprise), []).append((ligne, str(entreprise)))
    # Lignes reprises par leur nom exact : jamais prises par un nom approché.
    prises = {existantes[cle(d.offre.entreprise, d.offre.intitule)] for d in a_ecrire
              if cle(d.offre.entreprise, d.offre.intitule) in existantes}

    ajoutees = maj = 0
    for d in sorted(dossiers, key=lambda d: d.envoi.horodatage or datetime.max):
        if not a_exporter(d):
            continue
        champs = valeurs(d)
        ligne = existantes.get(cle(d.offre.entreprise, d.offre.intitule))
        if ligne is None and "poste" not in colonnes:
            # « Nordex » saisi à la main = « NORDEX FRANCE (NORDEX FRANCE SAS) », dans le même département
            ligne = coeurs.get((_coeur(d.offre.entreprise), _departement(d.offre.lieu)))
            if ligne in prises:
                ligne = None
        if ligne is not None:
            prises.add(ligne)
        nouvelle = ligne is None
        if nouvelle:
            derniere += 1
            ligne = derniere
            existantes[cle(d.offre.entreprise, d.offre.intitule)] = ligne
            if ligne - 1 > ligne_entete:  # style de la ligne précédente
                for col in range(1, ws.max_column + 1):
                    source, cible = ws.cell(ligne - 1, col), ws.cell(ligne, col)
                    if source.has_style:
                        cible._style = copy(source._style)
            ajoutees += 1
        else:
            maj += 1
        # Valeurs brutes écrites par une version précédente (nom en double, code NAF, code postal) : remplacées.
        brutes = {"entreprise": d.offre.entreprise, "lieu": d.offre.lieu, "secteur": d.fiche_entreprise.secteur or d.offre.secteur}
        for champ, col in colonnes.items():
            cellule = ws.cell(ligne, col)
            if champ in brutes and not nouvelle and cellule.value == brutes[champ] and champs[champ]:
                cellule.value = champs[champ]
                continue
            if champ == "statut":
                etat = _etat(d)
                valeur = _choisir(_liste_deroulante(ws, cellule), MOTS_STATUT[etat], LIBELLES_STATUT[etat])
            elif (champ in DESCRIPTIFS and not nouvelle and cellule.value not in (None, "")
                  and not (champ == "presentiel" and str(cellule.value).startswith("Oui : "))):
                continue  # saisi à la main (nom, notes…) : intouchable ; « Oui : » vient de l'outil
            else:
                valeur = champs[champ]
            if valeur in (None, ""):
                continue  # une case déjà remplie à la main n'est pas vidée
            cellule.value = valeur
            if champ == "presentiel":
                _surligner(cellule)
            if isinstance(valeur, (date, datetime)) and cellule.number_format == "General":
                cellule.number_format = "DD/MM/YYYY HH:MM" if isinstance(valeur, datetime) else "DD/MM/YYYY"
        coeur = _coeur(d.offre.entreprise)
        autres = [(n, nom) for n, nom in homonymes.get(coeur, []) if n != ligne] if coeur else []
        if autres and "remarques" in colonnes and ws.cell(ligne, colonnes["remarques"]).value in (None, ""):
            n, nom = autres[0]  # même entreprise ailleurs dans le tableau (autre site) : à toi de voir
            ws.cell(ligne, colonnes["remarques"]).value = f"Même entreprise que la ligne {n} ({nom}) ?"

    departements = {}  # ligne -> département écrit dans le tableau
    if "lieu" in colonnes:
        departements = {n: _departement(ws.cell(n, colonnes["lieu"]).value) for n in range(ligne_entete + 1, derniere + 1)}
    for dm in demarches:
        ligne = existantes.get(cle(dm.entreprise, ""))
        if ligne is not None and _departement(dm.lieu) and departements.get(ligne) not in ("", None, _departement(dm.lieu)):
            ligne = None  # même nom, autre site (« ENGIE » du Nord n'est pas celui de la Somme)
        if ligne is None and "poste" not in colonnes:
            ligne = coeurs.get((_coeur(dm.entreprise), _departement(dm.lieu)))
        nouvelle = ligne is None
        if nouvelle:
            derniere += 1
            ligne = derniere
            existantes[cle(dm.entreprise, "")] = ligne
            if ligne - 1 > ligne_entete:
                for col in range(1, ws.max_column + 1):
                    source, cible = ws.cell(ligne - 1, col), ws.cell(ligne, col)
                    if source.has_style:
                        cible._style = copy(source._style)
            ajoutees += 1
        else:
            maj += 1
        champs = valeurs_demarche(dm)
        for champ, col in colonnes.items():
            cellule = ws.cell(ligne, col)
            if not nouvelle and cellule.value not in (None, ""):
                continue  # le tableau est à toi : une démarche ne remplit que les cases vides
            if champ == "statut":
                mots, libelle = STATUTS_DEMARCHE[dm.statut]
                liste = _liste_deroulante(ws, cellule)
                valeur = _choisir(liste, mots, "") if liste else libelle
            else:
                valeur = champs.get(champ)
            if valeur in (None, ""):
                continue
            cellule.value = valeur
            if isinstance(valeur, date) and cellule.number_format == "General":
                cellule.number_format = "DD/MM/YYYY"

    for table in ws.tables.values():  # tableau Excel mis en forme : on l'étend aux nouvelles lignes
        debut, fin = table.ref.split(":")
        fin_col = re.match(r"[A-Z]+", fin).group(0)
        if int(re.search(r"\d+", fin).group(0)) < derniere:
            table.ref = f"{debut}:{fin_col}{derniere}"

    temporaire = chemin.with_name(f"~rs_{chemin.name}")
    try:
        classeur.save(temporaire)
        os.replace(temporaire, chemin)
    except PermissionError as e:
        temporaire.unlink(missing_ok=True)
        raise ErreurExcel(f"{chemin.name} est ouvert dans Excel : le fermer, il sera mis à jour au prochain passage") from e
    return ajoutees, maj
