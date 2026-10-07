"""Tenue du tableau Excel de suivi de l'utilisateur."""

from datetime import date, datetime

import pytest
from conftest import offre
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table

from recherche_stage.excel import ErreurExcel, mettre_a_jour, reconnaitre_colonnes
from recherche_stage.hachage import id_dossier
from recherche_stage.modeles import Dossier, Statut


def dossier(statut=Statut.envoyee, **champs) -> Dossier:
    o = offre(**champs)
    d = Dossier(id=id_dossier(o.entreprise, o.intitule), offre=o, statut=statut)
    d.envoi.canal, d.envoi.destinataire = "email", o.contact_email
    d.envoi.horodatage = datetime(2026, 10, 6, 9, 12)
    return d


@pytest.fixture
def classeur(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Candidatures"
    ws["A1"] = "Suivi de mes candidatures de stage"
    ws.append([])
    ws.append(["Nom de l'entreprise", "Intitulé du poste", "Ville", "Date d'envoi", "Moyen", "Statut",
               "Date de relance", "Réponse", "Mes notes", "Dates du stage"])
    ws.append(["Exemple Énergie SAS", "Stage maintenance industrielle H/F", "Dunkerque", None, None, "En attente",
               None, None, "appelé le RH", "déc.–janv."])
    ws.append(["Boîte à moi", "Stage bureau d'études", "Lille", date(2026, 9, 1), "Mail", "Refusée", None, None, "", None])
    ws["A5"].font = Font(bold=True)
    statuts = DataValidation(type="list", formula1='"En attente,Relancée,Entretien obtenu,Refusée,Non retenu,Acceptée"')
    ws.add_data_validation(statuts)
    statuts.add("F4:F200")
    ws.add_table(Table(displayName="Suivi", ref="A3:J5"))
    chemin = tmp_path / "Suivi_candidatures_stage_1.xlsx"
    wb.save(chemin)
    return chemin


def test_colonnes_reconnues_a_leur_intitule():
    colonnes = reconnaitre_colonnes({1: "Nom de l'entreprise", 2: "Intitulé du poste", 3: "Date d'envoi", 4: "Dates du stage",
                                     5: "Date de relance", 6: "Lien de l'offre", 7: "Contact RH", 8: "Divers"})
    assert colonnes == {"entreprise": 1, "poste": 2, "date_candidature": 3, "relance": 5, "lien": 6, "contact": 7}


def test_mise_a_jour_sans_abimer_le_tableau(classeur):
    existante = dossier()  # même offre que la ligne saisie à la main (forme juridique en plus)
    nouvelle = dossier(entreprise="Autre Énergie", contact_email="rh@autre.fr")
    nouvelle.statut = Statut.entretien
    nouvelle.suivi.reponse = "entretien"
    a_deposer = dossier(Statut.en_attente_humain, entreprise="Sans Adresse", contact_email="")
    a_deposer.etape_attente, a_deposer.envoi.canal, a_deposer.envoi.horodatage = "depot", "formulaire", None
    ignoree = dossier(Statut.abandonnee, entreprise="Rejetée")

    assert mettre_a_jour(classeur, [existante, nouvelle, a_deposer, ignoree]) == (2, 1)

    ws = load_workbook(classeur)["Candidatures"]
    assert ws["A4"].value == "Exemple Énergie SAS"  # nom saisi à la main gardé
    assert ws["D4"].value.date() == date(2026, 10, 6) and ws["E4"].value == "Mail"
    assert ws["F4"].value == "En attente"  # valeur de la liste déroulante
    assert ws["I4"].value == "appelé le RH" and ws["J4"].value == "déc.–janv."  # notes et colonne inconnue intactes
    assert [ws.cell(5, c).value for c in (1, 6)] == ["Boîte à moi", "Refusée"]  # ligne sans dossier intacte
    lignes = {ws.cell(r, 1).value: r for r in range(6, ws.max_row + 1)}
    assert ws.cell(lignes["Autre Énergie"], 6).value == "Entretien obtenu" and ws.cell(lignes["Autre Énergie"], 8).value == "Entretien"
    assert ws.cell(lignes["Sans Adresse"], 6).value == "À déposer" and ws.cell(lignes["Sans Adresse"], 4).value is None
    assert ws.cell(6, 1).font.bold  # style de la ligne précédente repris
    assert ws.tables["Suivi"].ref == "A3:J7"
    assert classeur.with_name("Suivi_candidatures_stage_1.avant_rs.xlsx").exists()

    assert mettre_a_jour(classeur, [existante, nouvelle, a_deposer]) == (0, 3)  # pas de doublon au passage suivant


def test_fichier_introuvable_ou_sans_entreprise(tmp_path):
    with pytest.raises(ErreurExcel, match="introuvable"):
        mettre_a_jour(tmp_path / "absent.xlsx", [])
    wb = Workbook()
    wb.active.append(["Divers", "Autre"])
    wb.save(tmp_path / "vide.xlsx")
    with pytest.raises(ErreurExcel, match="Entreprise"):
        mettre_a_jour(tmp_path / "vide.xlsx", [])


def test_cycle_et_validation_tiennent_le_tableau(ctx, classeur):
    import random

    from conftest import FauxExpediteur, FauxSource, evaluation_critique, evaluation_offre, sortie_redaction
    from test_orchestrateur import fiche_extraite

    from recherche_stage.orchestrateur import Orchestrateur

    ctx.config.suivi.excel = str(classeur)
    ctx.llm.recherche = ("x", ["https://exemple-energie.fr/actus"])
    ctx.llm.reponses.update({"scoring": evaluation_offre(), "enquete": fiche_extraite(), "redaction": sortie_redaction(),
                             "critique": evaluation_critique()})
    affiche = []
    o = Orchestrateur(ctx, [FauxSource("jobteaser", [offre()])], FauxExpediteur(), rng=random.Random(0), dormir=lambda s: None)
    o.progres = affiche.append
    o.cycle()
    assert "Tableau Excel : 0 ligne(s) ajoutée(s), 0 mise(s) à jour" in affiche  # rien de validé encore
    [d] = ctx.store.lister()
    o.valider(d.id)
    o.cycle(collecter=False)
    ws = load_workbook(classeur)["Candidatures"]
    assert ws["D4"].value is not None and ws["E4"].value == "Mail"

    ctx.config.suivi.excel = str(classeur.with_name("absent.xlsx"))
    assert "non mis à jour" in o.mettre_a_jour_excel()  # signalé, sans faire tomber le cycle


def test_tableau_a_liste_a_envoyer_noms_lieux_et_activite(tmp_path):
    """Tableau avec « À envoyer » dans la liste, nom écrit à la main, et lignes
    écrites par une ancienne version (code NAF, code postal, nom en double)."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Suivi candidatures"])
    ws.append([])
    ws.append(["Entreprise", "Activité", "Lieu", "Contact", "Date d'envoi", "Statut", "Notes"])
    ws.append(["Turbinor", "Constructeur éolien", "Berry-au-Bac (02)", None, None, "À envoyer", None])
    ws.append(["EAU CLAIRE (EAU CLAIRE)", "71.12B", "SAINT-QUENTIN (02100)", "rh@eau.fr", date(2026, 10, 6), "À envoyer", None])
    statuts = DataValidation(type="list", formula1='"À envoyer,Envoyée,Relancée,Entretien,Acceptée,Refusée"')
    ws.add_data_validation(statuts)
    statuts.add("F4:F60")
    chemin = tmp_path / "suivi.xlsx"
    wb.save(chemin)

    envoyee = dossier(entreprise="EAU CLAIRE (EAU CLAIRE)", lieu="SAINT-QUENTIN (02100)", contact_email="rh@eau.fr")
    envoyee.fiche_entreprise.secteur = "71.12B"
    meme = dossier(Statut.validee, entreprise="TURBINOR FRANCE (TURBINOR FRANCE SAS)", lieu="LAON (02000)")
    autre_site = dossier(Statut.validee, entreprise="TURBINOR FRANCE SAS", lieu="VILLENEUVE-D'ASCQ (59650)")
    autre_site.offre.secteur = "28.11Z"
    assert mettre_a_jour(chemin, [envoyee, meme, autre_site]) == (1, 2)

    ws = load_workbook(chemin).active
    assert [c.value for c in ws[5]][:3] == ["EAU CLAIRE", "Ingénierie, études techniques", "Saint-Quentin (02)"]
    assert ws["F5"].value == "Envoyée"  # et non « À envoyer »
    assert ws["A4"].value == "Turbinor" and ws["C4"].value == "Berry-au-Bac (02)"  # saisi à la main : intact
    assert ws["F4"].value == "À envoyer"
    assert [c.value for c in ws[6]][:3] == ["TURBINOR FRANCE SAS", "Fabrication de moteurs et turbines", "Villeneuve-d'Ascq (59)"]
    assert ws["G6"].value == "Même entreprise que la ligne 4 (Turbinor) ?"
