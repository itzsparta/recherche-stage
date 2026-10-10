"""Démarches faites hors de l'outil (forum, recommandation) : lecture du fichier,
tableau Excel, pas de candidature spontanée en double, rappels de relance."""

import random
from datetime import date, datetime

import pytest
from conftest import FauxExpediteur, FauxHttp, evaluation_critique, evaluation_offre, lettre, offre, sortie_redaction
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from recherche_stage.agents.scoring import AgentScoring
from recherche_stage.demarches import Demarche, ErreurDemarches, charger, correspondance
from recherche_stage.excel import mettre_a_jour
from recherche_stage.hachage import id_dossier
from recherche_stage.interface import Application
from recherche_stage.modeles import Dossier, Statut
from recherche_stage.orchestrateur import Orchestrateur
from recherche_stage.presentiel import distance_km

FICHIER = """
demarches:
  - entreprise: Dillinger France
    statut: envoyee
    date: 2026-09-28
    canal: Recommandation d'un salarié
    lieu: Grande-Synthe (59)
  - entreprise: STEF
    statut: a_faire
    date: 2026-10-08
    canal: Forum AgoraJob
    notes: plusieurs sites en Hauts-de-France, se présenter directement
  - entreprise: Assystem
    statut: en_attente
    date: 2026-10-08
    canal: Forum AgoraJob
    lieu: Beauvais (60)
    notes: alternance l'an prochain
    relance: 2026-10-08
  - entreprise: Engie / Enedis
    date: 2026-10-08
    canal: Forum AgoraJob
    notes: stand visité
"""


@pytest.fixture
def demarches(ctx, tmp_path):
    chemin = tmp_path / "demarches.yaml"
    chemin.write_text(FICHIER, encoding="utf-8")
    ctx.config.suivi.demarches = str(chemin)
    return chemin


def test_lecture_et_correspondance(demarches, tmp_path):
    liste = charger(demarches)
    assert [d.statut for d in liste] == ["envoyee", "a_faire", "en_attente", "contact"]
    assert correspondance(liste, "STEF TRANSPORT LILLE", "Lesquin (59)").entreprise == "STEF"  # tous ses sites
    assert correspondance(liste, "DILLINGER FRANCE (DILLINGER FRANCE SAS)", "DUNKERQUE (59640)").entreprise == "Dillinger France"
    assert correspondance(liste, "Dillinger France", "Metz (57)") is None  # autre département
    assert correspondance(liste, "ENEDIS ENGIE") and correspondance(liste, "ENGIE") is None  # tous les mots du nom
    assert charger(tmp_path / "absent.yaml") == []
    (tmp_path / "faux.yaml").write_text("demarches:\n  - entreprise: X\n    statut: envoyé\n", encoding="utf-8")
    with pytest.raises(ErreurDemarches, match="démarche X : statut"):
        charger(tmp_path / "faux.yaml")


def test_pas_de_spontanee_en_double_mais_offres_en_tete(ctx, demarches):
    ctx.llm.reponses["scoring"] = evaluation_offre()
    spontanee = offre(entreprise="STEF TRANSPORT LILLE", intitule="Candidature spontanée", type="spontanee", lieu="Lesquin (59)")
    offre_stef = offre(entreprise="STEF", intitule="Stage logistique entrepôt", lieu="Lesquin (59)")
    rencontree = offre(entreprise="ENEDIS ENGIE", intitule="Candidature spontanée", type="spontanee")
    dossiers = [Dossier(id=id_dossier(o.entreprise, o.intitule), offre=o) for o in (spontanee, offre_stef, rencontree)]
    AgentScoring(ctx).executer(dossiers)
    bloque, offre_gardee, contact = dossiers
    assert bloque.statut == Statut.abandonnee and bloque.motif.startswith("deja_en_contact : STEF, Forum AgoraJob le 08/10/2026")
    assert offre_gardee.statut == Statut.scoree and offre_gardee.prioritaire  # une offre précise reste proposée, en tête
    assert contact.statut == Statut.scoree and contact.prioritaire  # stand visité : rien de bloqué


def test_tableau_excel_et_rappel_de_relance(ctx, demarches, tmp_path, horloge):
    wb = Workbook()
    ws = wb.active
    ws.append(["Suivi candidatures"])
    ws.append([])
    ws.append(["Entreprise", "Activité", "Lieu", "Contact", "Date d'envoi", "Statut", "Relance le", "Notes"])
    ws.append(["Dillinger France", "Tôles fortes", "Grande-Synthe (59)", "Recommandé par un employé", date(2026, 9, 28),
               "Entretien", None, "Relancer via l'employé"])
    ws.append(["Assystem", None, "Lille (59)", None, None, "À envoyer", None, None])  # même nom, autre site
    for n in range(5, 12):
        ws.cell(n, 7, f'=IF(E{n}="","",E{n}+10)')
    liste = DataValidation(type="list", formula1='"À envoyer,Envoyée,Relancée,Entretien,Acceptée,Refusée"')
    ws.add_data_validation(liste)
    liste.add("F4:F60")
    chemin = tmp_path / "suivi.xlsx"
    wb.save(chemin)

    assert mettre_a_jour(chemin, [], demarches=charger(demarches)) == (3, 1)  # la ligne Assystem du Nord : pas touchée
    assert mettre_a_jour(chemin, [], demarches=charger(demarches)) == (0, 4)  # retrouvées : pas de ligne en double
    ws = load_workbook(chemin).active
    lignes = {ws.cell(n, 1).value: [c.value for c in ws[n]] for n in range(4, ws.max_row + 1) if ws.cell(n, 1).value}
    assert lignes["Dillinger France"][5] == "Entretien" and lignes["Dillinger France"][7] == "Relancer via l'employé"  # à toi
    assert lignes["STEF"][5] == "À envoyer" and lignes["STEF"][7].startswith("Forum AgoraJob (08/10/2026) : plusieurs sites")
    assert lignes["STEF"][4] is None and lignes["STEF"][6].startswith("=IF")  # pas d'envoi : ni date, ni relance inventée
    assert [ws.cell(n, 6).value for n in range(4, ws.max_row + 1) if ws.cell(n, 1).value == "Assystem"] == ["À envoyer", None]
    assert lignes["Assystem"][5] is None and lignes["Assystem"][6].date() == date(2026, 10, 8)  # « En attente » absent de la liste
    assert lignes["Engie / Enedis"][7] == "Forum AgoraJob (08/10/2026) : stand visité"

    o = Orchestrateur(ctx, [], FauxExpediteur(), rng=random.Random(0), dormir=lambda s: None)
    assert o.rappeler_demarches() == 1 and o.rappeler_demarches() == 0  # STEF « à faire », une fois ; relance pas encore due
    a_faire = next(e for e in ctx.store.escalades_ouvertes() if e["type"] == "a_faire:STEF")
    assert "se présenter directement" in a_faire["question"] and "sur_place: true" in a_faire["question"]
    horloge.instant = datetime(2026, 10, 9, 9, 0)
    assert o.rappeler_demarches() == 1 and o.rappeler_demarches() == 0  # un seul rappel par date
    rappel = next(e for e in ctx.store.escalades_ouvertes() if e["type"] == "relance:Assystem")
    assert "Relance prévue le 08/10/2026 : Assystem (Forum AgoraJob" in rappel["question"]
    assert Demarche(entreprise="X", statut="contact").bloque_spontanee is False


class Annuaire(FauxHttp):
    """Annuaire des entreprises : sociétés trouvées selon le département demandé."""

    def __init__(self, par_departement: dict):
        super().__init__()
        self.par_departement = par_departement

    def get_json(self, url, params=None, entetes=None):
        self.requetes.append((url, params))
        return {"results": self.par_departement.get((params or {}).get("departement"), [])}


def societe(nom, commune, code_postal, point, etat="A"):
    return {"nom_complet": nom, "siren": str(abs(hash(nom)) % 10**9), "activite_principale": "52.10A",
            "tranche_effectif_salarie": "41",  # 500 salariés et plus
            "matching_etablissements": [{"siret": f"{abs(hash(nom)) % 10**14}", "code_postal": code_postal,
                                         "libelle_commune": commune, "adresse": f"ZI {commune} {code_postal}",
                                         "latitude": str(point[0]), "longitude": str(point[1]), "etat_administratif": etat}]}


LILLE, LESQUIN, ARRAS = (50.6292, 3.0573), (50.589, 3.111), (50.291, 2.777)


def test_se_presenter_sur_place_prepare_la_candidature_de_la_tournee(ctx, tmp_path):
    chemin = tmp_path / "demarches.yaml"
    chemin.write_text("""
demarches:
  - {entreprise: STEF, statut: a_faire, date: 2026-10-08, canal: Forum AgoraJob, sur_place: true,
     notes: "Plusieurs sites en Hauts-de-France ; se présenter directement sur un site"}
  - {entreprise: Roquette, statut: a_faire, date: 2026-10-08, canal: Forum AgoraJob,
     notes: Suivre les offres sur leur site carrière}
""", encoding="utf-8")
    ctx.config.suivi.demarches = str(chemin)
    ctx.config.presentiel.domicile = "Lille"
    ctx.store.definir_parametre("geo:lille", list(LILLE))
    ctx.config.sources.annuaire_departements = ["59", "62"]
    ctx.http = Annuaire({
        "59": [societe("STEFANI MATERIAUX", "LILLE", "59000", LILLE),  # pas STEF : écartée
               societe("STEF TRANSPORT LILLE", "LESQUIN", "59810", LESQUIN)],
        "62": [societe("STEF LOGISTIQUE ARRAS", "ARRAS", "62000", ARRAS),
               societe("STEF TRANSPORT CALAIS", "CALAIS", "62100", (50.95, 1.86), etat="F")],  # site fermé
    })
    stef = "STEF TRANSPORT LILLE"
    ctx.config.llm.recherche_web = False
    ctx.llm.reponses.update({
        "enquete": {"activite": "Transport frigorifique", "activite_source": "", "faits": [], "contact_email": "",
                    "contact_source": "", "fiable": True},
        "redaction": sortie_redaction(lettre(entreprise=stef, intitule="Stage maintenance (candidature spontanée)"), mail=(
            f"Bonjour,\n\nVeuillez trouver ci-joint ma candidature spontanée chez {stef} : mon CV et ma lettre de motivation. "
            "Mon stage en maintenance préventive et ma pratique de la GMAO pourraient servir vos équipes.\n\n"
            "Bien cordialement,\nCamille Martin\ncamille.martin@exemple.fr\n06 12 34 56 78")),
        "critique": evaluation_critique(),
    })
    o = Orchestrateur(ctx, [], FauxExpediteur(), rng=random.Random(0), dormir=lambda s: None)
    o.apprentissage.choisir_variante = lambda: "experience_terrain"
    rapport = o.cycle(envoyer=False, suivre=False)
    assert rapport["depots_demarches"] == 1 and rapport["enquetes"] == 1

    [d] = ctx.store.lister()
    assert (d.offre.entreprise, d.offre.source, d.offre.adresse, d.prioritaire) == (stef, "demarche", "ZI LESQUIN 59810", True)
    # Grand groupe et score bas sur le secteur : préparée quand même, l'entreprise a demandé de passer.
    assert d.presentiel.conseille and d.presentiel.raison.startswith("se présenter sur place, d'après ta démarche")
    assert d.score.justification.startswith("À déposer sur place")
    assert d.envoi.canal == "main_propre" and d.statut == Statut.en_attente_humain and d.etape_attente == "envoi"
    d.fiche_entreprise.contact = "rh@stef.exemple"
    assert o.envoi.choisir_canal(d) == ("main_propre", "ZI LESQUIN 59810")  # jamais par mail, même avec une adresse

    alertes = {e["type"]: e["question"] for e in ctx.store.escalades_ouvertes() if not e["dossier_id"]}
    km = distance_km(LILLE, LESQUIN), distance_km(LILLE, ARRAS)
    assert f"le plus proche, {stef}, ZI LESQUIN 59810, à {km[0]:.0f} km" in alertes["sur_place:STEF"]
    assert f"Autres sites : STEF LOGISTIQUE ARRAS (ARRAS (62000), {km[1]:.0f} km)." in alertes["sur_place:STEF"]
    assert "CALAIS" not in alertes["sur_place:STEF"] and "a_faire:STEF" not in alertes
    assert "rs detecter-ats" in alertes["a_faire:Roquette"]  # suivre leurs offres : comment les faire arriver ici

    # Dans l'interface : sur la carte avant validation, puis arrêt de la tournée, déposée en main propre.
    app = Application(lambda: (ctx, o), ouvrir=lambda chemin: None, arriere_plan=False)
    [repere] = app.etat()["tournee"]["reperes"]
    assert (repere["entreprise"], repere["etat"], repere["point"]) == (stef, "a_valider", LESQUIN)
    app.action({"action": "valider", "dossier": d.id})
    carte = next(e for e in app.etat()["file"] if e["dossier"] and e["dossier"]["id"] == d.id)
    assert carte["dossier"]["etape"] == "depot" and carte["actions"] == ["deposer", "rejeter", "ouvrir"]  # ni robot ni site
    assert [a["entreprise"] for a in app.etat()["tournee"]["arrets"]] == [stef]
    app.action({"action": "deposer", "dossier": d.id, "preuve": "en main propre, 12/10/2026"})
    d = ctx.store.get(d.id)
    assert d.statut == Statut.envoyee and d.envoi.canal == "main_propre" and d.presentiel.depose_le  # cochée dans la tournée

    requetes = len(ctx.http.requetes)
    assert o.preparer_depots_demarches() == 0 and len(ctx.http.requetes) == requetes  # une seule fois par démarche


def test_se_presenter_sur_place_adresse_donnee_ou_site_introuvable(ctx, tmp_path):
    chemin = tmp_path / "demarches.yaml"
    chemin.write_text("""
demarches:
  - {entreprise: Exemple Froid, statut: a_faire, sur_place: true, adresse: "1 rue du Port, 59140 Dunkerque"}
  - {entreprise: Inconnue Logistique, statut: a_faire, sur_place: true}
""", encoding="utf-8")
    ctx.config.suivi.demarches = str(chemin)
    ctx.config.presentiel.domicile = "Lille"
    ctx.store.definir_parametre("geo:lille", list(LILLE))
    ctx.store.definir_parametre("geo:1 rue du port, 59140 dunkerque", [51.03, 2.37])
    ctx.http = Annuaire({})
    o = Orchestrateur(ctx, [], FauxExpediteur(), rng=random.Random(0), dormir=lambda s: None)
    assert o.preparer_depots_demarches() == 1
    [d] = ctx.store.lister()
    assert (d.offre.entreprise, d.offre.latitude, d.presentiel.distance_km) == ("Exemple Froid", 51.03, round(distance_km(LILLE, (51.03, 2.37)), 1))
    assert [p for _, p in ctx.http.requetes] == [{"q": "Inconnue Logistique", "per_page": 25, "etat_administratif": "A"}]
    [alerte] = [e for e in ctx.store.escalades_ouvertes() if e["type"] == "sur_place:Inconnue Logistique"]
    assert "aucun de ses sites trouvé" in alerte["question"] and "adresse:" in alerte["question"]
    assert o.preparer_depots_demarches() == 0 and len(ctx.http.requetes) == 1  # pas recherché à chaque cycle

    # Le site est ajouté au fichier : nouvel essai.
    chemin.write_text(chemin.read_text(encoding="utf-8").replace("sur_place: true}", 'sur_place: true, adresse: "2 quai Est, 59000 Lille"}'),
                      encoding="utf-8")
    assert o.preparer_depots_demarches() == 1


def test_revue_depot_en_main_propre(ctx):
    from recherche_stage.modeles import Presentiel
    from recherche_stage.revue import Revue

    o = Orchestrateur(ctx, [], FauxExpediteur(), rng=random.Random(0), dormir=lambda s: None)
    site = "ZI LESQUIN 59810"
    of = offre(entreprise="STEF TRANSPORT LILLE", intitule="Stage (candidature spontanée)", type="spontanee",
               source="demarche", adresse=site)
    d = Dossier(id=id_dossier(of.entreprise, of.intitule), offre=of, statut=Statut.en_attente_humain, etape_attente="depot")
    d.envoi.canal, d.envoi.destinataire = "main_propre", site
    d.presentiel = Presentiel(conseille=True, raison="se présenter sur place, d'après ta démarche", adresse=site)
    ctx.store.save(d)
    ctx.store.escalader(d.id, "depot_manuel", "Déposer la candidature en main propre")
    entrees, sortie, ouverts = ["o", "d"], [], []
    bilan = Revue(ctx, o, demander=lambda invite: entrees.pop(0), afficher=sortie.append, ouvrir=ouverts.append).executer()
    assert bilan["deposees"] == 1 and f"À apporter en main propre chez STEF TRANSPORT LILLE : {site}" in "\n".join(sortie)
    assert len(ouverts) == 1  # le dossier seulement : pas de recherche de page carrières
    d = ctx.store.get(d.id)
    assert d.statut == Statut.envoyee and d.presentiel.depose_le and d.envoi.canal == "main_propre"
