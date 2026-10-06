"""Signalement des candidatures à déposer en main propre."""

import random

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table

from conftest import FauxHttp, offre
from recherche_stage.config import ParamsPresentiel
from recherche_stage.hachage import id_dossier
from recherche_stage.modeles import Dossier, Statut
from recherche_stage.orchestrateur import Orchestrateur
from recherche_stage.presentiel import API_ADRESSE, distance_km, evaluer, geocoder
from recherche_stage.reseau import ErreurReseau
from recherche_stage.sources.annuaire import SourceAnnuaire

BOULOGNE = (50.7264, 1.6147)
DESVRES = (50.6670, 1.8330)  # ~17 km
LAIGNEVILLE = (49.2900, 2.4400)  # ~170 km


@pytest.fixture
def orchestrateur(ctx):
    return Orchestrateur(ctx, [], rng=random.Random(0), dormir=lambda s: None)


def ban(point):
    lat, lon = point
    return {"features": [{"geometry": {"coordinates": [lon, lat]}}]}


def spontanee(statut=Statut.validee, point=DESVRES, tranche="11", **champs) -> Dossier:
    base = dict(entreprise="Atelier Côte d'Opale", intitule="Stage ingénieur (candidature spontanée)", type="spontanee",
                source="annuaire", contact_email=None, lieu="DESVRES (62240)", adresse="12 RUE DU PORT 62240 DESVRES",
                latitude=point[0] if point else None, longitude=point[1] if point else None, tranche_effectif=tranche)
    base.update(champs)
    o = offre(**base)
    return Dossier(id=id_dossier(o.entreprise, o.intitule), offre=o, statut=statut)


def test_distance():
    assert 14 < distance_km(BOULOGNE, DESVRES) < 20
    assert 150 < distance_km(BOULOGNE, LAIGNEVILLE) < 190


def test_evaluer():
    params = ParamsPresentiel()
    p = evaluer(spontanee(), BOULOGNE, DESVRES, params)
    assert p.conseille and "à 17 km" in p.raison and "moins de 50 salariés" in p.raison
    assert "pas d'adresse mail publique" in p.raison and p.adresse.endswith("DESVRES")
    assert not evaluer(spontanee(), BOULOGNE, LAIGNEVILLE, params).conseille  # trop loin
    grand = evaluer(spontanee(tranche="42"), BOULOGNE, DESVRES, params)  # 1000 salariés et plus
    assert not grand.conseille and "grande entreprise" in grand.raison
    assert not evaluer(spontanee(type="offre"), BOULOGNE, DESVRES, params).conseille  # offre publiée
    # Ancien dossier sans le champ : la tranche est relue dans la description.
    ancien = spontanee(tranche=None, description="Activité principale (NAF) : 28.11Z. Tranche d'effectif INSEE : 53")
    assert not evaluer(ancien, BOULOGNE, DESVRES, params).conseille


def test_geocoder_commune_et_cache(ctx):
    http = FauxHttp({API_ADRESSE: ban(DESVRES)})
    assert geocoder(http, ctx.store, "DESVRES (62240)") == DESVRES
    assert http.requetes[0][1] == {"q": "DESVRES", "limit": 1, "postcode": "62240", "type": "municipality"}
    assert geocoder(http, ctx.store, "DESVRES  (62240)") == DESVRES and len(http.requetes) == 1  # cache

    class Panne(FauxHttp):
        def get_json(self, url, params=None, entetes=None):
            raise ErreurReseau("hors ligne")

    assert geocoder(Panne(), ctx.store, "Calais") is None
    assert ctx.store.parametre("geo:calais") is None  # panne : pas mise en cache


def test_signale_dans_l_interface_et_l_excel(ctx, orchestrateur, tmp_path):
    ctx.http.json_par_url[API_ADRESSE] = ban(BOULOGNE)  # domicile : la ville du CV
    proche, loin = spontanee(), spontanee(point=LAIGNEVILLE, entreprise="Loin SAS")
    for d in (proche, loin):
        ctx.store.save(d)
    wb = Workbook()
    ws = wb.active
    ws.append(["Entreprise", "Poste", "Statut"])
    ws.append(["Déjà là", "Stage", "Envoyée"])
    ws.add_table(Table(displayName="Suivi", ref="A1:C2"))
    chemin = tmp_path / "suivi.xlsx"
    wb.save(chemin)
    ctx.config.suivi.excel = str(chemin)

    orchestrateur.mettre_a_jour_excel()

    assert ctx.store.get(proche.id).presentiel.conseille
    assert not ctx.store.get(loin.id).presentiel.conseille
    ws = load_workbook(chemin).active
    assert ws["D1"].value == "Dépôt en présentiel"
    lignes = {ws.cell(r, 1).value: ws.cell(r, 4) for r in range(2, ws.max_row + 1)}
    cellule = lignes["Atelier Côte d'Opale"]
    assert cellule.value.startswith("Oui : à 17 km") and "12 RUE DU PORT" in cellule.value
    assert cellule.fill.start_color.rgb.endswith("C6EFCE")
    assert lignes["Loin SAS"].value is None
    assert ws.tables["Suivi"].ref == "A1:D4" and ws.tables["Suivi"].tableColumns[-1].name == "Dépôt en présentiel"

    from recherche_stage.interface import Application

    carte = Application(lambda: (ctx, orchestrateur), ouvrir=lambda chemin: None)._dossier(ctx.store.get(proche.id))
    assert 15 < carte["presentiel"]["distance_km"] < 20 and carte["presentiel"]["adresse"].startswith("12 RUE")


def test_sans_domicile_rien_n_est_juge(ctx, orchestrateur):
    ctx.store.save(spontanee())
    affiche = []
    orchestrateur.progres = affiche.append
    assert orchestrateur.evaluer_presentiel() == 0  # la fausse API ne trouve pas la ville
    assert any("presentiel.domicile" in m for m in affiche)
    assert ctx.store.lister()[0].presentiel is None


def test_annuaire_garde_les_coordonnees_de_l_etablissement(ctx):
    source = SourceAnnuaire(FauxHttp(), ["28.11Z"], ["62"], ctx.store)
    o = source._vers_offre({
        "nom_complet": "Atelier", "siren": "123456789", "tranche_effectif_salarie": "11",
        "siege": {"libelle_commune": "PARIS", "code_postal": "75001"},
        "matching_etablissements": [{"libelle_commune": "DESVRES", "code_postal": "62240", "adresse": "12 RUE DU PORT 62240 DESVRES",
                                     "latitude": "50.667", "longitude": "1.833"}],
    }, "Stage ingénieur", "62")
    assert (o.latitude, o.longitude, o.tranche_effectif) == (50.667, 1.833, "11") and o.adresse.startswith("12 RUE")


def test_ordre_de_tournee_itineraires_et_excel_apres_depot():
    from datetime import datetime

    from recherche_stage.modeles import Presentiel
    from recherche_stage.presentiel import itineraires, ordonner, texte_excel

    assert ordonner(BOULOGNE, [("loin", LAIGNEVILLE), ("inconnu", None), ("proche", DESVRES)]) == ["proche", "loin", "inconnu"]
    liens = itineraires("Calais", [f"arret {i}" for i in range(12)])
    assert len(liens) == 2 and liens[0].startswith("https://www.google.com/maps/dir/Calais/arret%200/")
    assert liens[1].startswith("https://www.google.com/maps/dir/arret%208/arret%209/")  # repart du dernier arrêt
    p = Presentiel(conseille=True, raison="à 5 km", adresse="1 rue X", depose_le=datetime(2026, 12, 8))
    assert texte_excel(p) == "Fait : déposée le 08/12/2026 — 1 rue X"
